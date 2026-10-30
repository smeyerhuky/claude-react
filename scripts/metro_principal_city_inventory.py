"""Bind official principal-city memberships to all qualifying CBSA research domains.

No coordinates, graph endpoints, access evidence or anchors are adopted.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT/'docs/research/principal-city-inventory'
SOURCE = REFERENCE/'list2_2023.xlsx'
DOMAINS = ROOT/'docs/research/utah-parent-reference/domains.json'
SOURCE_SHA256 = '1e1091de28d1ad1423cebf18e7cdb6db4fbf2204fd357bc68c24e8261dcce161'
SOURCE_URL = 'https://www2.census.gov/programs-surveys/metro-micro/geographies/reference-files/2023/delineation-files/list2_2023.xlsx'
STATE_FIPS = dict(pair.split(':') for pair in ('01:AL 02:AK 04:AZ 05:AR 06:CA 08:CO 09:CT 10:DE 11:DC 12:FL 13:GA 15:HI 16:ID 17:IL 18:IN 19:IA 20:KS 21:KY 22:LA 23:ME 24:MD 25:MA 26:MI 27:MN 28:MS 29:MO 30:MT 31:NE 32:NV 33:NH 34:NJ 35:NM 36:NY 37:NC 38:ND 39:OH 40:OK 41:OR 42:PA 44:RI 45:SC 46:SD 47:TN 48:TX 49:UT 50:VT 51:VA 53:WA 54:WV 55:WI 56:WY 72:PR').split())
HEADERS = ('CBSA Code', 'CBSA Title', 'Metropolitan/Micropolitan Statistical Area',
           'Principal City Name', 'FIPS State Code', 'FIPS Place Code')


def require(condition, message):
    if not condition: raise ValueError(message)


def memberships(rows):
    result = {}; seen = set()
    for row in rows:
        if all(value is None for value in row): continue
        # Explicit Census footer text is not a city row. Reject malformed codes
        # elsewhere rather than heuristically normalizing names or leading zeros.
        if isinstance(row[0], str) and row[0].startswith(('Note:', 'Source:', 'Internet Release Date:')):
            continue
        require(len(row) == 6 and isinstance(row[0], str) and re.fullmatch('[0-9]{5}', row[0]),
                'Unexpected non-record row in official principal-city table')
        cbsa, title, kind, name, state, place = row
        require(isinstance(title, str) and title and isinstance(name, str) and name
                and kind in ('Metropolitan Statistical Area', 'Micropolitan Statistical Area')
                and isinstance(state, str) and re.fullmatch('[0-9]{2}', state)
                and isinstance(place, str) and re.fullmatch('[0-9]{5}', place),
                'Explicit source title, city and zero-preserving FIPS codes required')
        identity = (cbsa, state+place)
        require(identity not in seen, 'Duplicate principal-city membership')
        seen.add(identity)
        record = result.setdefault(cbsa, dict(title=title, kind=kind, cities=[]))
        require(record['title'] == title and record['kind'] == kind, 'Conflicting CBSA source identity')
        record['cities'].append(dict(principal_city_id=state+place, name=name,
                                    state_fips=state, place_fips=place))
    return result


def project(source=SOURCE, domains_path=DOMAINS):
    raw = source.read_bytes()
    require(len(raw) == 79368 and hashlib.sha256(raw).hexdigest() == SOURCE_SHA256,
            'Exact original principal-city workbook required')
    workbook = load_workbook(source, read_only=True, data_only=True)
    require(workbook.sheetnames == ['List 2'], 'Review changed workbook sheets')
    rows = list(workbook.active.values); workbook.close()
    require(rows[1][0] == 'List 2. PRINCIPAL CITIES OF METROPOLITAN AND MICROPOLITAN STATISTICAL AREAS, JULY 2023'
            and rows[2] == HEADERS, 'Review changed principal-city table/vintage')
    by_cbsa = memberships(rows[3:])
    domain_bytes = domains_path.read_bytes(); reference = json.loads(domain_bytes)
    domains = reference['domains']; ids = [d['cbsa'] for d in domains]
    require(len(ids) == len(set(ids)) == 367 and sorted(ids) == reference['source_qualifying_cbsa_ids'],
            'Complete unique qualifying CBSA research domain required')
    records = []
    for domain in sorted(domains, key=lambda d: d['cbsa']):
        cbsa = domain['cbsa']
        require(cbsa in by_cbsa and by_cbsa[cbsa]['kind'] == 'Metropolitan Statistical Area',
                'Missing qualifying metro membership; never omit a CBSA')
        entry = by_cbsa[cbsa]
        records.append(dict(cbsa=cbsa, qualifying_domain_name=domain['name'],
                            source_cbsa_title=entry['title'], source_population=domain['source_population'],
                            included_in_existing_csv=domain['included_in_existing_csv'],
                            principal_cities=sorted(entry['cities'], key=lambda c: c['principal_city_id']),
                            anchor_selection_status='not_selected', adopted_anchor_ids=[], access_verified=False,
                            release_eligible=False))
    # This grouping follows each source city's state, not the metro's native
    # representative point. A cross-state CBSA can occur in several groups.
    state_groups = []
    require(all(c['state_fips'] in STATE_FIPS for r in records for c in r['principal_cities']),
            'Principal-city jurisdiction outside explicit state/DC/PR scope')
    for fips, abbreviation in sorted(STATE_FIPS.items()):
        cities = [dict(cbsa=r['cbsa'], **c) for r in records for c in r['principal_cities'] if c['state_fips'] == fips]
        state_groups.append(dict(state_fips=fips, state=abbreviation, principal_city_memberships=cities,
                                 principal_city_membership_count=len(cities),
                                 cbsa_ids=sorted({c['cbsa'] for c in cities}),
                                 anchor_selection_status='not_selected', release_eligible=False))
    return dict(schema='metro-principal-city-inventory-v1',
                source=dict(url=SOURCE_URL, sha256=SOURCE_SHA256, bytes=len(raw), delineation='July 2023',
                            retrieved_utc='2026-10-30', workbook_sheet='List 2'),
                domains_sha256=hashlib.sha256(domain_bytes).hexdigest(),
                qualifying_cbsa_count=len(records), principal_city_membership_count=sum(len(r['principal_cities']) for r in records),
                qualifying_cbsa_ids=sorted(ids),
                qualifiers_absent_from_existing_csv=[r['cbsa'] for r in records if not r['included_in_existing_csv']],
                states=state_groups, jurisdiction_scope='50 states, DC and Puerto Rico; principal-city membership only, not geographic coverage',
                metros=records, anchors_adopted=False, access_verified=False, release_eligible=False,
                warning='Official July 2023 principal-city membership joined by exact CBSA ID to the retained qualifying domain. Not current vehicle entry evidence, an anchor selection, coordinate registry or approval. Source vintage and any later changes require review; no title splitting or population-based city filtering.')


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('--write', action='store_true'); args = p.parse_args()
    result = project(); target = REFERENCE/'report.json'
    if args.write: target.write_text(json.dumps(result, indent=2, sort_keys=True)+'\n')
    else: require(json.loads(target.read_text()) == result, 'Retained principal-city inventory differs')
    print(json.dumps({k: result[k] for k in ('qualifying_cbsa_count', 'principal_city_membership_count', 'anchors_adopted', 'release_eligible')}, indent=2))
