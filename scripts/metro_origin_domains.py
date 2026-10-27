#!/usr/bin/env python3
"""Research source inventory/domains; preserve the public CONUS dataset."""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import zipfile

from shapely.geometry import Point

from colorado_tiger2025_geography import _native_crs, _read_shapes


def population_records(rows, threshold=100000):
    if type(threshold) is not int or threshold <= 0:
        raise ValueError("Invalid population threshold")
    header = next((i for i, row in enumerate(rows[:20]) if any("geographic area" in str(c).lower() for c in row)), None)
    if header is None:
        raise ValueError("Missing Census population header")
    area = next(i for i, c in enumerate(rows[header]) if "geographic area" in str(c).lower())
    years = [i for row in rows[max(0, header-2):header+4] for i, c in enumerate(row) if str(c).strip() == "2025"]
    if not years:
        raise ValueError("Missing 2025 population column")
    year = max(years)
    result, seen = [], set()
    for row in rows[header+1:]:
        name = str(row[area] or "").strip().lstrip(".")
        if "metro area" not in name.lower():
            continue
        raw = row[year] if year < len(row) else None
        try:
            value = int(str(raw).replace(",", ""))
        except (ValueError, TypeError) as error:
            raise ValueError("Missing/noninteger metropolitan population") from error
        if isinstance(raw, bool) or value < 0 or name in seen:
            raise ValueError("Invalid or duplicate metropolitan population record")
        seen.add(name)
        if value >= threshold:
            result.append({"name": name, "population": value})
    if not result:
        raise ValueError("Empty qualifying source population inventory")
    return result


def build_domains(existing_rows, populations, geometries, registry=None):
    existing = {}
    for row in existing_rows:
        if int(row["population"]) < 100000:
            continue
        if row["cbsa"] in existing:
            raise ValueError("Duplicate existing CBSA")
        existing[row["cbsa"]] = row
    if not existing:
        raise ValueError("Missing existing inventory")
    by_name, seen_ids = {}, set()
    for fields, geometry in geometries:
        cbsa = fields["GEOID"]
        if cbsa in seen_ids:
            raise ValueError("Duplicate source CBSA geography")
        seen_ids.add(cbsa)
        if fields["LSAD"] != "M1":
            continue
        if (fields["NAMELSAD"] in by_name or geometry.is_empty or not geometry.is_valid or
                geometry.geom_type not in ("Polygon", "MultiPolygon")):
            raise ValueError("Invalid or ambiguous metropolitan geography")
        by_name[fields["NAMELSAD"]] = (fields, geometry)
    joined, source_names, membership = {}, set(), []
    urban_rows = {}
    if registry is not None:
        if registry["release_eligible"] is not False or registry["access_certified"] is not False:
            raise ValueError("Research scenario safety mismatch")
        rows = registry["scenarios"]["colorado-named-city-point-leads-v1"]
        urban_rows = {r["cbsa"]: r for r in rows}
        if len(urban_rows) != len(rows) or set(urban_rows) != set(existing):
            raise ValueError("Scenario/current CSV inventory mismatch")
    for pop in populations:
        name = pop["name"]
        if name in source_names or name not in by_name:
            raise ValueError("Duplicate or unmatched exact Census metropolitan name: "+name)
        source_names.add(name)
        fields, geometry = by_name[name]
        cbsa = fields["GEOID"]
        if cbsa in joined:
            raise ValueError("Duplicate source population-to-CBSA join")
        point = [float(fields["INTPTLON"]), float(fields["INTPTLAT"])]
        if not geometry.covers(Point(*point)):
            raise ValueError("Source internal point outside corresponding native polygon")
        row = {"cbsa": cbsa, "name": name, "source_population": pop["population"],
               "native_representative_lonlat": point, "native_bbox_lonlat": list(geometry.bounds),
               "included_in_existing_csv": cbsa in existing, "representative_point_inside_native_polygon": True,
               "access_verified": False, "road_origin_verified": False}
        if cbsa in existing:
            old = existing[cbsa]
            if old["name"] != name or int(old["population"]) != pop["population"]:
                raise ValueError("Existing/source population identity changed; review scope before reuse")
        joined[cbsa] = row
        urban = urban_rows.get(cbsa)
        if urban and urban["origin_lonlat"] is not None:
            membership.append({"cbsa": cbsa, "origin_lonlat": urban["origin_lonlat"],
                               "inside_native_cbsa_polygon": bool(geometry.covers(Point(*urban["origin_lonlat"]))),
                               "access_verified": False, "datum_alignment_verified": False})
    if set(existing)-set(joined):
        raise ValueError("Existing qualifying source missing from national-source domains")
    extras = sorted(set(joined)-set(existing))
    return {"release_eligible": False, "access_certified": False, "complete_global_competition": False,
            "road_coverage_certified": False, "datum_alignment_verified": False, "native_crs": "EPSG:4269",
            "source_table_scope": "2025 Census metropolitan population table for United States and Puerto Rico",
            "existing_csv_cbsa_ids": sorted(existing), "source_qualifying_cbsa_ids": sorted(joined),
            "source_qualifiers_absent_from_existing_csv": extras,
            "domains": [joined[k] for k in sorted(joined)], "selected_urban_lead_membership": membership,
            "urban_domain_hypothesis": "Any future approved urban anchor must be inside its corresponding sourced CBSA footprint; this is an unapproved research constraint, not an anchor or official place crosswalk",
            "warning": "Native CBSA bounding boxes are origin-domain diagnostics, never road-graph coverage certificates. Preserve existing public CONUS CSV. No source excluded; native-to-OSM datum alignment and full costed connectors remain unverified."}


def main():
    from openpyxl import load_workbook
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--metros", type=Path, required=True)
    p.add_argument("--population", type=Path, required=True)
    p.add_argument("--cbsa-zip", type=Path, required=True)
    p.add_argument("--scenarios", type=Path)
    p.add_argument("--output", type=Path, required=True)
    args = p.parse_args()
    lock_path = Path(__file__).with_name("metro_origin_domain_sources.json")
    lock = json.loads(lock_path.read_text())
    source_hashes = {}
    for kind, path in (("population", args.population), ("geography", args.cbsa_zip)):
        if path.stat().st_size > (5000000 if kind == "population" else 130000000):
            raise ValueError("Source exceeds bounded archive size")
        source_hashes[kind] = hashlib.sha256(path.read_bytes()).hexdigest()
        if source_hashes[kind] != lock[kind+"_sha256"]:
            raise ValueError("Source bytes changed; review and version the source lock")
    workbook = load_workbook(args.population, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError("Review added workbook sheets before claiming complete source coverage")
    populations = population_records(list(workbook.active.values))
    with args.metros.open(newline="", encoding="utf-8") as handle:
        existing = list(csv.DictReader(handle))
    registry = json.loads(args.scenarios.read_text()) if args.scenarios else None
    csv_hash = hashlib.sha256(args.metros.read_bytes()).hexdigest()
    if registry is not None and registry["source_provenance"]["metros_csv_sha256"] != csv_hash:
        raise ValueError("Scenario source CSV mismatch")
    result = build_domains(existing, populations, _read_shapes(args.cbsa_zip, 5000), registry)
    with zipfile.ZipFile(args.cbsa_zip) as source:
        result["source_projection"] = _native_crs(source, args.cbsa_zip.stem)
    result["source_provenance"] = {**lock, "observed_sha256": source_hashes, "metros_csv_sha256": csv_hash,
                                   "workbook_sheet": workbook.active.title}
    if args.scenarios:
        result["source_provenance"]["scenario_registry_sha256"] = hashlib.sha256(args.scenarios.read_bytes()).hexdigest()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True)+"\n")
    print(json.dumps({"existing_qualifiers": len(result["existing_csv_cbsa_ids"]), "source_qualifiers": len(result["source_qualifying_cbsa_ids"]),
                      "source_qualifiers_absent_from_existing_csv": result["source_qualifiers_absent_from_existing_csv"], "release_eligible": False}, indent=2))


if __name__ == "__main__":
    main()
